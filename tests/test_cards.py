import openpyxl
import pytest

from accesspath.cards import load_library, LibraryError
from conftest import DATA


def edited(tmp_path, card_id, **changes):
    """Copy of the real library with one card row changed."""
    wb = openpyxl.load_workbook(DATA / "card_library.xlsx")
    ws = wb["Cards"]
    head = [c.value for c in ws[1]]
    for row in ws.iter_rows(min_row=2):
        if row[0].value == card_id:
            for col, v in changes.items():
                row[head.index(col)].value = v
    path = tmp_path / "lib.xlsx"
    wb.save(path)
    return path


def test_real_library_loads(library):
    assert len(library.cards) == 38
    assert {c.status for c in library.cards.values()} == {"sourced", "interpolated", "placeholder"}
    assert "stroke_R_moderate" in library.profiles


def test_one_row_edit_changes_the_number(tmp_path):
    lib = load_library(edited(tmp_path, "base-elbow_R-flex", ROM_max_deg=140))
    assert lib.cards["base-elbow_R-flex"].ROM_max_deg == 140


def test_dash_in_number_column_is_explained(tmp_path):
    with pytest.raises(LibraryError, match="leave the cell empty"):
        load_library(edited(tmp_path, "base-elbow_R-flex", ROM_max_deg="—"))


def test_placeholder_with_numbers_is_rejected(tmp_path):
    with pytest.raises(LibraryError, match="placeholder but has numbers"):
        load_library(edited(tmp_path, "arthritis-shoulder-abd", ROM_max_deg=120))


def test_sourced_needs_a_citation(tmp_path):
    with pytest.raises(LibraryError, match="no citation"):
        load_library(edited(tmp_path, "base-elbow_R-flex", source="NEEDED: not yet sourced"))


def test_filling_a_placeholder(tmp_path):
    lib = load_library(edited(tmp_path, "arthritis-shoulder-abd", ROM_max_deg=120,
                              status="sourced", source="Example et al. 2027"))
    assert lib.cards["arthritis-shoulder-abd"].ROM_max_deg == 120
