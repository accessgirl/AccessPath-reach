import openpyxl
import pytest

from accesspath.cards import load_library, LibraryError
from accesspath.privacy import check_tree, check_fields, IdentifyingDataError
from accesspath.verify import run, load_json
from conftest import DATA


@pytest.mark.parametrize("key", ["client_name", "Date of Birth", "home address", "Phone", "email", "DOB", "SSN",
                                 "diagnosis", "patient-name"])
def test_identifying_fields_are_rejected(key):
    with pytest.raises(IdentifyingDataError, match="client vault"):
        check_fields([key], "test")


@pytest.mark.parametrize("key", ["profile_id", "posture", "clear_width_m", "position", "wall_normal", "card_ids",
                                 "body_band", "fixture", "room_id", "description"])
def test_engine_fields_pass(key):
    check_fields([key], "test")


def test_room_with_an_address_is_rejected(library):
    room = load_json(DATA / "rooms" / "test_bathroom.json")
    room["site"] = {"address": "12 Oak St"}
    with pytest.raises(IdentifyingDataError, match="room test_bathroom.site"):
        run(library, room, load_json(DATA / "tasks.json")["tasks"], DATA / "body.json", ["baseline_wheelchair"])


def test_real_room_and_library_carry_no_identity(library):
    check_tree(load_json(DATA / "rooms" / "test_bathroom.json"), "room")


def test_profiles_sheet_with_a_name_column_is_rejected(tmp_path):
    wb = openpyxl.load_workbook(DATA / "card_library.xlsx")
    ws = wb["Profiles"]
    ws.cell(1, ws.max_column + 1, "client_name")
    path = tmp_path / "lib.xlsx"
    wb.save(path)
    with pytest.raises(LibraryError, match="client_name"):
        load_library(path)
