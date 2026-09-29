"""Card library: load the spreadsheet (the source of truth) and check it before anything uses it."""
from dataclasses import dataclass, field, asdict
import json

import openpyxl

from .isncsci import parse_sci_name

STATUSES = {"sourced", "interpolated", "placeholder"}
RANGE_KINDS = {"joint_limits", "score_band", "task_threshold", "sweep_width", "envelope_shift", "reach_reduction"}
POSTURES = {"standing", "seated_wheelchair"}
SIDES = {"R", "L", "both", "none"}


class LibraryError(ValueError):
    """The spreadsheet has a problem that would make results wrong. The message names the sheet and row."""


@dataclass
class Card:
    card_id: str
    joint_id: str
    parent_joint: str
    DOF_axis: str
    condition: str
    severity: str
    ROM_min_deg: float | None
    ROM_max_deg: float | None
    range_kind: str
    strength: str | None
    control: str | None
    laterality: str | None
    dependency_flags: str | None
    notes: str | None
    source: str | None
    status: str
    ISB_term: str | None = None

    @property
    def is_baseline(self):
        return self.condition == "baseline"


@dataclass
class Profile:
    profile_id: str
    description: str
    posture: str
    mobility_aid: str
    affected_side: str
    card_ids: list[str] = field(default_factory=list)


@dataclass
class PopulationCap:
    cap_id: str
    applies_to_posture: str
    rule: str
    value: float
    notes: str | None
    source: str | None
    status: str


@dataclass
class Library:
    cards: dict[str, Card]
    profiles: dict[str, Profile]
    caps: list[PopulationCap]

    def to_json(self):
        return json.dumps({
            "cards": [asdict(c) for c in self.cards.values()],
            "profiles": [asdict(p) for p in self.profiles.values()],
            "population_caps": [asdict(c) for c in self.caps],
        }, indent=2)


def _rows(wb, sheet):
    if sheet not in wb.sheetnames:
        raise LibraryError(f"The spreadsheet has no '{sheet}' sheet.")
    rows = list(wb[sheet].iter_rows(values_only=True))
    header = [str(h).strip() if h is not None else "" for h in rows[0]]
    for n, row in enumerate(rows[1:], start=2):
        if all(v is None or str(v).strip() == "" for v in row):
            continue
        yield n, {h: _clean(v) for h, v in zip(header, row) if h}


def _clean(v):
    if isinstance(v, str):
        v = v.strip()
        return v or None
    return v


def _number(v, where, col):
    if v is None:
        return None
    if isinstance(v, (int, float)):
        return float(v)
    try:
        return float(v)
    except ValueError:
        raise LibraryError(f"{where}: {col} is '{v}'. Use a number, or leave the cell empty if it isn't sourced yet.")


def _need(row, cols, where):
    missing = [c for c in cols if c not in row]
    if missing:
        raise LibraryError(f"{where}: missing column(s) {', '.join(missing)}.")


def load_library(path) -> Library:
    wb = openpyxl.load_workbook(path, data_only=True)
    cards: dict[str, Card] = {}
    for n, row in _rows(wb, "Cards"):
        where = f"Cards row {n}"
        _need(row, ["card_id", "joint_id", "DOF_axis", "condition", "ROM_min_deg", "ROM_max_deg",
                    "range_kind", "status", "source"], where)
        cid = row["card_id"]
        if not cid:
            raise LibraryError(f"{where}: card_id is empty.")
        if cid in cards:
            raise LibraryError(f"{where}: card_id '{cid}' is used twice.")
        card = Card(
            card_id=cid, joint_id=row["joint_id"], parent_joint=row.get("parent_joint"),
            DOF_axis=row["DOF_axis"], condition=row["condition"], severity=row.get("severity"),
            ROM_min_deg=_number(row["ROM_min_deg"], where, "ROM_min_deg"),
            ROM_max_deg=_number(row["ROM_max_deg"], where, "ROM_max_deg"),
            range_kind=row["range_kind"], strength=row.get("strength"), control=row.get("control"),
            laterality=row.get("laterality"), dependency_flags=row.get("dependency_flags"),
            notes=row.get("notes"), source=row.get("source"), status=row["status"],
            ISB_term=row.get("ISB_term"),
        )
        _check_card(card, where)
        cards[cid] = card

    profiles: dict[str, Profile] = {}
    for n, row in _rows(wb, "Profiles"):
        where = f"Profiles row {n}"
        _need(row, ["profile_id", "posture", "mobility_aid", "affected_side", "card_ids"], where)
        ids = [c.strip() for c in (row["card_ids"] or "").split(",") if c.strip()]
        p = Profile(row["profile_id"], row.get("description") or "", row["posture"],
                    row["mobility_aid"], str(row["affected_side"]), ids)
        if p.profile_id in profiles:
            raise LibraryError(f"{where}: profile_id '{p.profile_id}' is used twice.")
        if p.profile_id.upper().startswith("SCI"):
            try:
                parse_sci_name(p.profile_id)
            except ValueError as err:
                raise LibraryError(f"{where}: {err}")
        if p.posture not in POSTURES:
            raise LibraryError(f"{where}: posture '{p.posture}' should be one of {sorted(POSTURES)}.")
        if p.affected_side not in SIDES:
            raise LibraryError(f"{where}: affected_side '{p.affected_side}' should be one of {sorted(SIDES)}.")
        for cid in ids:
            if cid not in cards:
                raise LibraryError(f"{where}: card_id '{cid}' is not on the Cards sheet.")
            if cards[cid].is_baseline:
                raise LibraryError(f"{where}: '{cid}' is a baseline card; baseline cards are always applied, list only condition cards.")
        if ids and p.affected_side == "none":
            raise LibraryError(f"{where}: lists condition cards but affected_side is 'none'.")
        profiles[p.profile_id] = p

    caps = []
    if "Population Caps" in wb.sheetnames:
        for n, row in _rows(wb, "Population Caps"):
            where = f"Population Caps row {n}"
            caps.append(PopulationCap(row["cap_id"], row["applies_to_posture"], row["rule"],
                                      _number(row["value"], where, "value"), row.get("notes"),
                                      row.get("source"), row["status"]))
    return Library(cards, profiles, caps)


def _check_card(c: Card, where):
    if c.status not in STATUSES:
        raise LibraryError(f"{where}: status '{c.status}' should be one of {sorted(STATUSES)}.")
    if c.range_kind not in RANGE_KINDS:
        raise LibraryError(f"{where}: range_kind '{c.range_kind}' should be one of {sorted(RANGE_KINDS)}.")
    has_numbers = c.ROM_min_deg is not None or c.ROM_max_deg is not None
    if c.status == "placeholder" and has_numbers:
        raise LibraryError(f"{where}: '{c.card_id}' is a placeholder but has numbers. "
                           "Mark it sourced or interpolated, or empty the numbers.")
    if c.status == "sourced" and (not c.source or c.source.upper().startswith("NEEDED")):
        raise LibraryError(f"{where}: '{c.card_id}' is marked sourced but has no citation.")
    if c.status != "placeholder" and c.range_kind == "joint_limits" and c.ROM_max_deg is None:
        raise LibraryError(f"{where}: '{c.card_id}' is {c.status} but ROM_max_deg is empty.")
    for v, col in [(c.ROM_min_deg, "ROM_min_deg"), (c.ROM_max_deg, "ROM_max_deg")]:
        if v is not None and not 0 <= v <= 360:
            raise LibraryError(f"{where}: {col} {v} should be degrees from neutral between 0 and 360.")
    if c.ROM_min_deg is not None and c.ROM_max_deg is not None and c.ROM_min_deg > c.ROM_max_deg:
        raise LibraryError(f"{where}: ROM_min_deg is larger than ROM_max_deg.")
    if c.is_baseline and "_affected" in c.joint_id:
        raise LibraryError(f"{where}: baseline cards need a real side (_R / _L), not '_affected'.")
    if not c.is_baseline and "_affected" not in c.joint_id and c.joint_id != "trunk":
        raise LibraryError(f"{where}: condition cards use '<joint>_affected'; the profile picks the side.")
